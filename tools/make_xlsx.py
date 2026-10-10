"""data/problems.json -> 단원별 분류 엑셀(문항 분류 / 단원별 보기). 사용: python3 tools/make_xlsx.py out.xlsx"""
import json, sys, os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SEC_RANK = {"공통": 0, "확률과 통계": 1, "미적분": 2, "기하": 3}
d = json.load(open(os.path.join(ROOT, "data/problems.json"), encoding="utf-8"))
probs = sorted(d["problems"], key=lambda p: (p["year"], p["grade"], p["month"], p.get("form") or "", SEC_RANK[p["sec"]], p["n"]))
def suffix(p):   # 2021~ 3학년은 (선택과목), 2020 3학년은 (가형/나형)
    if p.get("form"): return f"({p['form']}형)"
    return f"({p['sec']})" if p["grade"] == 3 else ""
def gu(p): return p["form"] + "형" if p.get("form") else (p["sec"] if p["grade"] == 3 else None)
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor="1F3864"); HDR_FONT = Font(bold=True, color="FFFFFF")
wb = Workbook(); ws = wb.active; ws.title = "문항 분류"
heads = ["순번", "문항", "연도", "학년", "시행월", "문제번호", "구분", "과목", "단원", "성취기준 코드", "해당 성취기준", "문항 요지", "비고"]
ws.append(heads)
for i, p in enumerate(probs, 1):
    lab = f"{p['year']}년 {p['month']}월 {p['grade']}학년 {p['n']}번" + suffix(p) + (f" (={p['also']})" if p.get("also") else "")
    ws.append([i, lab, p["year"], p["grade"], f"{p['month']}월", p["n"], gu(p), p["subject"], p["unit"], p["code"], p["standard"], p["summary"], None])
# 비고 채우기 (분류 JSON / 6월 xlsx)
import openpyxl
notes = {}
x6 = openpyxl.load_workbook(os.path.join(ROOT, "data/classification_2021_06.xlsx"))["문항 분류"]
for r in x6.iter_rows(min_row=2, values_only=True):
    notes[(int(r[2]), int(r[3]), 6, None, r[6] or "공통", int(r[5]))] = r[12]
SEC_KEY = {"C": "공통", "P": "확률과 통계", "K": "미적분", "G": "기하"}
for f in os.listdir(os.path.join(ROOT, "data/classification")):
    parts = f[:-5].split("_"); y, g, m = parts[:3]; fm = {"ga": "가", "na": "나", "A": "A", "B": "B"}.get(parts[3]) if len(parts) > 3 else None
    dd = json.load(open(os.path.join(ROOT, "data/classification", f), encoding="utf-8"))
    for k, v in dd.items(): notes[(int(y), int(g), int(m), fm, SEC_KEY[k[0]], int(k[1:]))] = v.get("note") or None
for i, p in enumerate(probs, 2):
    ws.cell(i, 13).value = notes.get((p["year"], p["grade"], p["month"], p.get("form"), p["sec"], p["n"]))
widths = [6, 42, 7, 6, 7, 9, 12, 12, 22, 16, 55, 45, 40]
for c, w in zip("ABCDEFGHIJKLM", widths): ws.column_dimensions[c].width = w
for r, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row), 1):
    for c in row:
        c.border = BORDER
        if r == 1: c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            c.fill = PatternFill("solid", fgColor="F2F6FC" if r % 2 == 0 else "FFFFFF")
            c.alignment = Alignment(horizontal="left" if c.column >= 11 else "center", vertical="top", wrap_text=True)
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:M{ws.max_row}"

ws2 = wb.create_sheet("단원별 보기"); ws2.append(["과목", "단원", "문항 수", "문항 목록 (연도 > 학년 > 시행월 > 문제번호 순)"])
by = defaultdict(list)
for p in probs:
    lab = f"{p['year']} {p['grade']}학년 {p['month']}월 {p['n']}번" + suffix(p)
    by[(p["subject"], p["unit"])].append(lab)
order = [tuple(u) for u in d["units"]]
for k in by:
    if k not in order: order.append(k)
tot = 0
for k in order:
    v = by.get(k, []); tot += len(v)
    ws2.append([k[0], k[1], len(v), ", ".join(v) if v else None])
ws2.append(["합계", None, tot, None])
for c, w in zip("ABCD", [12, 24, 9, 110]): ws2.column_dimensions[c].width = w
for r, row in enumerate(ws2.iter_rows(min_row=1, max_row=ws2.max_row), 1):
    for c in row:
        c.border = BORDER
        if r == 1: c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left" if c.column == 4 else "center", vertical="top", wrap_text=True)
            if r == ws2.max_row: c.font = Font(bold=True)
ws2.freeze_panes = "A2"
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "data/classification_all.xlsx")
wb.save(out); print(out, len(probs), "rows")
